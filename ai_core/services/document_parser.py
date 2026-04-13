"""Document text extraction and chunking utilities.

Supports PDF, DOCX, and PPTX file formats. Extracted text is split into
overlapping chunks suitable for embedding and semantic search.
"""

import logging
from pathlib import Path

logger = logging.getLogger("ai_core")


def extract_text(file_path: str) -> str:
    """Extract plain text content from a document file.

    Args:
        file_path: Absolute path to the document file.

    Returns:
        Extracted text as a single string. Returns empty string
        if the file format is unsupported or extraction fails.
    """
    ext = Path(file_path).suffix.lower()

    try:
        if ext == ".pdf":
            return _extract_pdf(file_path)
        elif ext in (".doc", ".docx"):
            return _extract_docx(file_path)
        elif ext == ".pptx":
            return _extract_pptx(file_path)
        else:
            logger.warning("Unsupported file format for text extraction: %s", ext)
            return ""
    except Exception as e:
        logger.error("Failed to extract text from %s: %s", file_path, e)
        return ""


def chunk_text(
    text: str, chunk_size: int = 1000, overlap: int = 50
) -> list[str]:
    """Split text into overlapping chunks for embedding.

    Args:
        text: The full text to split.
        chunk_size: Maximum number of characters per chunk.
        overlap: Number of overlapping characters between consecutive chunks.

    Returns:
        A list of text chunks. Returns empty list if input text is blank.
    """
    if not text or not text.strip():
        return []

    chunks: list[str] = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = min(start + chunk_size, text_length)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start += chunk_size - overlap

    logger.info("Split text into %d chunks (size=%d, overlap=%d)", len(chunks), chunk_size, overlap)
    return chunks


# ---------------------------------------------------------------------------
# Private extraction helpers
# ---------------------------------------------------------------------------


def _extract_pdf(file_path: str) -> str:
    """Extract text from a PDF file using PyPDF2."""
    from PyPDF2 import PdfReader

    reader = PdfReader(file_path)
    pages: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            pages.append(page_text)

    text = "\n".join(pages)
    logger.info("Extracted %d characters from PDF (%d pages)", len(text), len(reader.pages))
    return text


def _extract_docx(file_path: str) -> str:
    """Extract text from a DOCX file using python-docx."""
    from docx import Document

    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    text = "\n".join(paragraphs)
    logger.info("Extracted %d characters from DOCX (%d paragraphs)", len(text), len(paragraphs))
    return text


def _extract_pptx(file_path: str) -> str:
    """Extract text from a PPTX file using python-pptx."""
    from pptx import Presentation

    prs = Presentation(file_path)
    slides_text: list[str] = []

    for slide in prs.slides:
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    paragraph_text = paragraph.text.strip()
                    if paragraph_text:
                        slide_parts.append(paragraph_text)
        if slide_parts:
            slides_text.append("\n".join(slide_parts))

    text = "\n\n".join(slides_text)
    logger.info("Extracted %d characters from PPTX (%d slides)", len(text), len(prs.slides))
    return text
